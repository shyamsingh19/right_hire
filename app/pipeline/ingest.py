from __future__ import annotations

import csv
import io
import logging
import re
from pathlib import Path

import httpx
import openpyxl

from app.config import settings

logger = logging.getLogger(__name__)

CANONICAL_FIELDS = ["name", "email", "resume_url", "yoe", "location"]

# Canonical column name mapping — covers common spreadsheet headers
_COL_MAP: dict[str, str] = {
    "name": "name",
    "full name": "name",
    "candidate name": "name",
    "applicant": "name",
    "applicant name": "name",
    "email": "email",
    "email address": "email",
    "e-mail": "email",
    "mail": "email",
    "contact email": "email",
    "resume_url": "resume_url",
    "resume url": "resume_url",
    "cv url": "resume_url",
    "resume link": "resume_url",
    "resume_drive_link": "resume_url",
    "drive link": "resume_url",
    "google drive link": "resume_url",
    "cv link": "resume_url",
    "portfolio": "resume_url",
    "attach your cv": "resume_url",
    "attach cv": "resume_url",
    "upload cv": "resume_url",
    "upload resume": "resume_url",
    "yoe": "yoe",
    "years of experience": "yoe",
    "experience": "yoe",
    "years": "yoe",
    "exp": "yoe",
    "total experience": "yoe",
    "location": "location",
    "city": "location",
    "city/country": "location",
    "country": "location",
    "address": "location",
    "region": "location",
    "place": "location",
}

# Keyword hints used for fuzzy fallback when exact match fails
_FIELD_KEYWORDS: dict[str, list[str]] = {
    "name": ["name", "candidate", "applicant"],
    "email": ["email", "mail", "e-mail"],
    "resume_url": ["resume", "cv", "portfolio", "link", "url", "attach", "drive", "upload"],
    "yoe": ["year", "exp", "yoe"],
    "location": ["location", "city", "country", "place", "address", "region"],
}


def detect_column_mapping(raw_headers: list[str]) -> dict[str, str | None]:
    """Return {canonical_field: raw_header_or_None} for each known field.

    First tries exact lookup in _COL_MAP, then falls back to keyword substring
    matching. A field is left None when no column matches.
    """
    mapping: dict[str, str | None] = {f: None for f in CANONICAL_FIELDS}
    claimed: set[str] = set()

    # Pass 1: exact matches
    for raw in raw_headers:
        canonical = _COL_MAP.get(raw.strip().lower())
        if canonical and mapping[canonical] is None and raw not in claimed:
            mapping[canonical] = raw
            claimed.add(raw)

    # Pass 2: keyword fuzzy for unmatched fields
    for field, keywords in _FIELD_KEYWORDS.items():
        if mapping[field] is not None:
            continue
        for raw in raw_headers:
            if raw in claimed:
                continue
            lower = raw.strip().lower()
            if any(kw in lower for kw in keywords):
                mapping[field] = raw
                claimed.add(raw)
                break

    return mapping


def _normalize_header(raw: str) -> str:
    return _COL_MAP.get(raw.strip().lower(), raw.strip().lower().replace(" ", "_"))


def apply_column_mapping(rows: list[dict], mapping: dict[str, str | None]) -> list[dict]:
    """Re-key rows using an explicit {canonical_field: raw_header} mapping."""
    reverse = {v: k for k, v in mapping.items() if v is not None}
    result = []
    for row in rows:
        new_row: dict = {}
        for raw_key, value in row.items():
            canonical = reverse.get(raw_key, raw_key)
            new_row[canonical] = value
        result.append(new_row)
    return result


def parse_excel(file_bytes: bytes) -> list[dict]:
    """Parse an .xlsx file and return a list of normalised row dicts."""
    wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    raw_headers = [str(h) if h is not None else "" for h in rows[0]]
    headers = [_normalize_header(h) for h in raw_headers]

    result: list[dict] = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue  # skip blank rows
        record: dict = {}
        for header, value in zip(headers, row):
            if not header:
                continue
            if value is not None:
                record[header] = value
        if record:
            result.append(record)

    wb.close()
    return result


def parse_excel_raw(file_bytes: bytes) -> tuple[list[str], list[dict]]:
    """Return (raw_headers, rows_keyed_by_raw_header) without normalising."""
    wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], []
    raw_headers = [str(h) if h is not None else "" for h in rows[0]]
    result = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        record = {h: (str(v) if v is not None else "") for h, v in zip(raw_headers, row)}
        if any(v for v in record.values()):
            result.append(record)
    return raw_headers, result


def parse_csv(file_bytes: bytes) -> list[dict]:
    """Parse a .csv file and return a list of normalised row dicts.

    Shares `_COL_MAP` / `_normalize_header` with `parse_excel` so both formats
    accept the same header aliases and produce identical row shapes.
    """
    text = file_bytes.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return []

    headers = [_normalize_header(h) for h in rows[0]]

    result: list[dict] = []
    for row in rows[1:]:
        if all(not v.strip() for v in row):
            continue  # skip blank rows
        record: dict = {}
        for header, value in zip(headers, row):
            if not header:
                continue
            value = value.strip()
            if value:
                record[header] = value
        if record:
            result.append(record)

    return result


def parse_csv_raw(file_bytes: bytes) -> tuple[list[str], list[dict]]:
    """Return (raw_headers, rows_keyed_by_raw_header) without normalising."""
    text = file_bytes.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return [], []
    raw_headers = [h.strip() for h in rows[0]]
    result = []
    for row in rows[1:]:
        if all(not v.strip() for v in row):
            continue
        record = {h: v.strip() for h, v in zip(raw_headers, row)}
        if any(v for v in record.values()):
            result.append(record)
    return raw_headers, result


def _is_drive_url(url: str) -> bool:
    return "drive.google.com" in url or "docs.google.com" in url


def _drive_to_direct(url: str) -> str:
    """Convert a Google Drive share URL to a direct download URL."""
    # e.g. https://drive.google.com/file/d/<ID>/view?...
    match = re.search(r"/file/d/([^/]+)", url)
    if match:
        file_id = match.group(1)
        return f"https://drive.google.com/uc?export=download&id={file_id}"
    # e.g. https://drive.google.com/open?id=<ID>
    match = re.search(r"id=([^&]+)", url)
    if match:
        file_id = match.group(1)
        return f"https://drive.google.com/uc?export=download&id={file_id}"
    return url


def fetch_drive_file(url: str, dest_path: str) -> str:
    """Download a file from Google Drive or a direct URL to *dest_path*.

    Returns the local path.
    """
    Path(dest_path).parent.mkdir(parents=True, exist_ok=True)

    if _is_drive_url(url):
        try:
            import gdown

            gdown.download(url, dest_path, quiet=True, fuzzy=True)
            return dest_path
        except Exception as exc:
            logger.warning("gdown failed (%s), falling back to direct download", exc)
            direct_url = _drive_to_direct(url)
            _http_download(direct_url, dest_path)
            return dest_path
    else:
        _http_download(url, dest_path)
        return dest_path


def _http_download(url: str, dest_path: str) -> None:
    with httpx.Client(timeout=60.0, follow_redirects=True) as client:
        with client.stream("GET", url) as resp:
            resp.raise_for_status()
            with open(dest_path, "wb") as fh:
                for chunk in resp.iter_bytes(chunk_size=8192):
                    fh.write(chunk)


def save_resume(candidate_id: str, content: bytes, ext: str) -> str:
    """Persist resume bytes under STORAGE_DIR and return the file path."""
    storage = Path(settings.storage_dir)
    storage.mkdir(parents=True, exist_ok=True)
    safe_ext = ext.lstrip(".").lower() or "pdf"
    file_path = storage / f"{candidate_id}.{safe_ext}"
    file_path.write_bytes(content)
    return str(file_path)
